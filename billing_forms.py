"""거래명세서 · 손실보전금 청구서 엑셀 서식."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet

import billing_database as billing

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
LABEL_FILL = PatternFill("solid", fgColor="EFEFEF")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
STAMP_FILL = PatternFill("solid", fgColor="F7F7F7")
WHITE = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(bold=True, color="333333", name="Malgun Gothic", size=10)
TITLE_FONT = Font(bold=True, name="Malgun Gothic", size=22)
SUB_FONT = Font(bold=True, name="Malgun Gothic", size=10)
BODY_FONT = Font(name="Malgun Gothic", size=9)
BOLD_FONT = Font(bold=True, name="Malgun Gothic", size=9)
MONEY_FONT = Font(bold=True, name="Malgun Gothic", size=12)
SMALL_FONT = Font(name="Malgun Gothic", size=8, color="555555")
THIN_SIDE = Side(style="thin", color="000000")
MED_SIDE = Side(style="medium", color="000000")
THIN = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
VERTICAL = Alignment(horizontal="center", vertical="center", wrap_text=True, textRotation=255)
LETTER_FONT = Font(name="Malgun Gothic", size=10.5)
LETTER_BOLD = Font(bold=True, name="Malgun Gothic", size=10.5)
LETTERHEAD_FONT = Font(bold=True, name="Malgun Gothic", size=14)
LAST_COL = 9
CLAIM_COLS = 8
REASON_PHRASES = {
    "자재지연": "원자재(부품) 입고 지연",
    "설계변경": "설계·사양 변경",
    "라인중단": "생산라인 중단 지시",
    "품질이슈": "품질 부적합(이슈)",
    "기타": "기타 귀책 사유",
}
ITEM_MIN_ROWS = 10
MONEY_FMT = "#,##0"

_DIGITS = ("", "일", "이", "삼", "사", "오", "육", "칠", "팔", "구")
_SMALL = ("", "십", "백", "천")
_BIG = ("", "만", "억", "조")


def amount_in_korean(amount: float) -> str:
    """수표·거래명세서용 한글 금액 (일금 ○○원정)."""
    n = int(round(float(amount or 0)))
    if n == 0:
        return "영원정"
    if n < 0:
        return "△" + amount_in_korean(-n)
    chunks: list[str] = []
    big_i = 0
    while n > 0:
        chunk = n % 10000
        if chunk:
            part = ""
            for unit in range(3, -1, -1):
                digit = (chunk // (10**unit)) % 10
                if digit:
                    part += _DIGITS[digit] + _SMALL[unit]
            chunks.append(part + _BIG[big_i])
        n //= 10000
        big_i += 1
    return "".join(reversed(chunks)) + "원정"


def _cell(ws: Worksheet, row: int, col: int, value: Any, *, fill=None, center: bool = False, bold: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, name="Malgun Gothic", size=10)
    cell.border = THIN
    cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
    if fill is not None:
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _a4_center(ws: Worksheet) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    try:
        ws.page_setup.horizontalCentered = True
        ws.page_setup.verticalCentered = True
    except Exception:
        pass
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    try:
        ws.sheet_view.showGridLines = False
    except Exception:
        pass
    try:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    except Exception:
        pass


def _widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _merge(ws: Worksheet, r1: int, c1: int, r2: int, c2: int) -> None:
    if r1 == r2 and c1 == c2:
        return
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _paint(
    ws: Worksheet,
    r1: int,
    c1: int,
    r2: int,
    c2: int,
    *,
    value: Any = None,
    font: Font | None = None,
    fill: PatternFill | None = None,
    align: Alignment | None = None,
    number_format: str | None = None,
) -> None:
    _merge(ws, r1, c1, r2, c2)
    for row in range(r1, r2 + 1):
        for col in range(c1, c2 + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN
            cell.font = font or BODY_FONT
            cell.alignment = align or CENTER
            if fill is not None:
                cell.fill = fill
            else:
                cell.fill = WHITE
    top = ws.cell(row=r1, column=c1, value=value)
    top.font = font or BODY_FONT
    top.alignment = align or CENTER
    if number_format:
        top.number_format = number_format
    if fill is not None:
        top.fill = fill


def _outline(ws: Worksheet, r1: int, c1: int, r2: int, c2: int) -> None:
    for row in range(r1, r2 + 1):
        for col in range(c1, c2 + 1):
            cell = ws.cell(row=row, column=col)
            left = MED_SIDE if col == c1 else THIN_SIDE
            right = MED_SIDE if col == c2 else THIN_SIDE
            top = MED_SIDE if row == r1 else THIN_SIDE
            bottom = MED_SIDE if row == r2 else THIN_SIDE
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def _parse_date(value: str) -> datetime:
    try:
        return datetime.strptime((value or "")[:10], "%Y-%m-%d")
    except ValueError:
        return datetime.now()


def _item_spec(row) -> str:
    spec = ""
    try:
        spec = str(row["item_spec"] or "")
    except (KeyError, IndexError, TypeError):
        spec = ""
    if spec.strip():
        return spec.strip()
    try:
        product_id = row["product_id"]
    except (KeyError, IndexError, TypeError):
        product_id = None
    return billing._item_spec_of(product_id, "")


def _write_party_box(
    ws: Worksheet,
    start_row: int,
    label_col: int,
    data_c1: int,
    data_c2: int,
    title: str,
    party: dict[str, str],
    *,
    stamp: bool = False,
) -> None:
    label_end = label_col
    _paint(
        ws, start_row, label_col, start_row + 4, label_end,
        value=title, font=BOLD_FONT, fill=LABEL_FILL, align=VERTICAL,
    )
    rows = (
        ("등록번호", party.get("biz_no", "")),
        ("상  호", party.get("company_name", "")),
        ("성  명", party.get("ceo_name", "")),
        ("주  소", party.get("address", "")),
        (
            "업태/종목",
            "\n".join(
                p
                for p in (
                    " / ".join(
                        x for x in (party.get("biz_type", ""), party.get("biz_item", "")) if x
                    ),
                    "  ".join(
                        x
                        for x in (
                            f"TEL {party['phone']}" if party.get("phone") else "",
                            f"FAX {party['fax']}" if party.get("fax") else "",
                        )
                        if x
                    ),
                )
                if p
            ),
        ),
    )
    for i, (label, value) in enumerate(rows):
        row = start_row + i
        if data_c2 - data_c1 >= 2:
            _paint(ws, row, data_c1, row, data_c1, value=label, font=BOLD_FONT, fill=LABEL_FILL)
            value_c1 = data_c1 + 1
        else:
            value_c1 = data_c1
        if stamp and i == 2 and data_c2 - value_c1 >= 1:
            _paint(ws, row, value_c1, row, data_c2 - 1, value=value, font=BODY_FONT, align=LEFT)
            _paint(ws, row, data_c2, start_row + 4, data_c2, value="(인)", font=SMALL_FONT, fill=STAMP_FILL)
        else:
            end_col = data_c2 - 1 if stamp and i >= 2 else data_c2
            if end_col < value_c1:
                end_col = value_c1
            _paint(ws, row, value_c1, row, end_col, value=value, font=BODY_FONT, align=LEFT)
    for i in range(5):
        ws.row_dimensions[start_row + i].height = 22


def _write_copy(
    ws: Worksheet,
    start_row: int,
    *,
    copy_label: str,
    doc_no: str,
    statement_date: str,
    seller: dict[str, str],
    buyer: dict[str, str],
    rows: list,
) -> int:
    issued = _parse_date(statement_date)
    date_text = f"{issued.year}년 {issued.month:02d}월 {issued.day:02d}일"
    supply_sum = vat_sum = total_sum = 0.0
    for row in rows:
        supply_sum += float(row["supply_price"] or 0)
        vat_sum += float(row["vat"] or 0)
        total_sum += float(row["total_amount"] or 0)

    r = start_row
    _paint(ws, r, 1, r, 4, value=f"No.  {doc_no}", font=BOLD_FONT, align=LEFT)
    _paint(ws, r, 5, r, 6, value="발행일", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(ws, r, 7, r, LAST_COL, value=date_text, font=BODY_FONT)
    ws.row_dimensions[r].height = 18

    r += 1
    _paint(ws, r, 1, r + 1, LAST_COL, value="거 래 명 세 서", font=TITLE_FONT)
    ws.row_dimensions[r].height = 26
    ws.row_dimensions[r + 1].height = 16
    r += 2
    _paint(ws, r, 1, r, LAST_COL, value=f"({copy_label})", font=SUB_FONT)
    ws.row_dimensions[r].height = 16

    r += 1
    party_top = r
    _write_party_box(ws, r, 1, 2, 5, "공급받는자", buyer)
    _write_party_box(ws, r, 6, 7, LAST_COL, "공급자", seller, stamp=True)
    r += 5

    _paint(ws, r, 1, r, 2, value="합계금액", font=BOLD_FONT, fill=TOTAL_FILL)
    _paint(
        ws, r, 3, r, 6,
        value=f"일금  {amount_in_korean(total_sum)}",
        font=BOLD_FONT, fill=TOTAL_FILL, align=LEFT,
    )
    _paint(
        ws, r, 7, r, LAST_COL,
        value=total_sum,
        font=MONEY_FONT, fill=TOTAL_FILL, align=RIGHT, number_format='"₩ "#,##0',
    )
    ws.row_dimensions[r].height = 22
    r += 1

    headers = ("월", "일", "품목", "규격", "수량", "단가", "공급가액", "세액", "비고")
    for col, heading in enumerate(headers, 1):
        _paint(ws, r, col, r, col, value=heading, font=BOLD_FONT, fill=HEADER_FILL)
    ws.row_dimensions[r].height = 18
    header_row = r
    r += 1

    line_count = max(ITEM_MIN_ROWS, len(rows))
    qty_sum = 0
    for i in range(line_count):
        values: list[Any] = ["", "", "", "", "", "", "", "", ""]
        if i < len(rows):
            row = rows[i]
            work = _parse_date(str(row["statement_date"]))
            qty_sum += int(row["quantity"] or 0)
            item_label = f"{row['item_name']}"
            code = str(row["item_code"] or "").strip()
            if code:
                item_label = f"{row['item_name']}\n({code})"
            values = [
                f"{work.month:02d}",
                f"{work.day:02d}",
                item_label,
                _item_spec(row),
                int(row["quantity"] or 0),
                float(row["unit_price"] or 0),
                float(row["supply_price"] or 0),
                float(row["vat"] or 0),
                row["remarks"] or "",
            ]
        for col, value in enumerate(values, 1):
            align = CENTER if col != 3 and col != 9 else (LEFT if col == 9 else CENTER)
            fmt = MONEY_FMT if col in (5, 6, 7, 8) and value != "" else None
            font = BODY_FONT
            _paint(ws, r, col, r, col, value=value if value != "" else None, font=font, align=align, number_format=fmt)
            if col == 3:
                ws.cell(row=r, column=col).alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
        ws.row_dimensions[r].height = 22 if i < len(rows) else 18
        r += 1

    _paint(ws, r, 1, r, 3, value="합  계", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, r, 4, r, 4, value="", fill=HEADER_FILL)
    _paint(ws, r, 5, r, 5, value=qty_sum, font=BOLD_FONT, fill=HEADER_FILL, number_format=MONEY_FMT)
    _paint(ws, r, 6, r, 6, value="", fill=HEADER_FILL)
    _paint(ws, r, 7, r, 7, value=supply_sum, font=BOLD_FONT, fill=HEADER_FILL, number_format=MONEY_FMT)
    _paint(ws, r, 8, r, 8, value=vat_sum, font=BOLD_FONT, fill=HEADER_FILL, number_format=MONEY_FMT)
    _paint(ws, r, 9, r, 9, value="", fill=HEADER_FILL)
    ws.row_dimensions[r].height = 20
    r += 1

    bank = " ".join(
        p for p in (
            seller.get("bank_name", ""),
            seller.get("bank_account", ""),
            f"예금주 {seller['bank_holder']}" if seller.get("bank_holder") else "",
        ) if p
    )
    manager = seller.get("manager_name", "")
    manager_phone = seller.get("manager_phone", "") or seller.get("phone", "")
    _paint(
        ws, r, 1, r, LAST_COL,
        value="위 금액을 정히 청구함.  인수 후 7일 이내 이의가 없으면 위 내용이 확인된 것으로 합니다.",
        font=BODY_FONT, align=LEFT,
    )
    ws.row_dimensions[r].height = 18
    r += 1
    _paint(
        ws, r, 1, r, 4,
        value=date_text,
        font=BOLD_FONT, align=CENTER,
    )
    _paint(ws, r, 5, r, 6, value=f"담당  {manager}", font=BODY_FONT, align=LEFT)
    _paint(ws, r, 7, r, LAST_COL, value=f"인수자  ____________", font=BODY_FONT, align=LEFT)
    ws.row_dimensions[r].height = 20
    r += 1
    _paint(ws, r, 1, r, 4, value=f"입금계좌  {bank}", font=BODY_FONT, align=LEFT)
    _paint(ws, r, 5, r, LAST_COL, value=f"연락처  {manager_phone}", font=BODY_FONT, align=LEFT)
    ws.row_dimensions[r].height = 18

    _outline(ws, start_row, 1, r, LAST_COL)
    _outline(ws, party_top, 1, party_top + 4, LAST_COL)
    _outline(ws, header_row, 1, r - 3, LAST_COL)
    return r


def export_statement(path: str | Path, rows: list, customer: str, statement_date: str) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "거래명세서"
    _a4_center(ws)
    seller = billing.get_company_profile()
    first = rows[0] if rows else None
    customer_id = None
    if first is not None:
        try:
            customer_id = first["customer_id"]
        except (KeyError, IndexError, TypeError):
            customer_id = None
    buyer = billing.party_for_customer(customer, customer_id)
    doc_no = billing.statement_doc_no(rows)

    last1 = _write_copy(
        ws, 1,
        copy_label="공급받는자 보관용",
        doc_no=doc_no,
        statement_date=statement_date,
        seller=seller,
        buyer=buyer,
        rows=rows,
    )
    cut = last1 + 1
    _paint(
        ws, cut, 1, cut, LAST_COL,
        value="- - - - -  절  취  선  - - - - -     (아래는 공급자 보관용)",
        font=SMALL_FONT,
    )
    ws.row_dimensions[cut].height = 14
    last2 = _write_copy(
        ws, cut + 1,
        copy_label="공급자 보관용",
        doc_no=doc_no,
        statement_date=statement_date,
        seller=seller,
        buyer=buyer,
        rows=rows,
    )
    _widths(ws, {1: 8, 2: 11, 3: 18, 4: 14, 5: 10, 6: 10, 7: 13, 8: 12, 9: 14})
    ws.print_area = f"A1:I{last2}"
    ws.page_setup.fitToHeight = 1
    wb.save(path)
    return path


def export_loss_claim(path: str | Path, row) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "업무협조전"
    _a4_center(ws)

    seller = billing.get_company_profile()
    buyer = billing.party_for_customer(str(row["customer_name"] or ""))
    issued = _parse_date(str(row["claim_date"]))
    date_text = f"{issued.year}년 {issued.month:02d}월 {issued.day:02d}일"
    doc_no = billing.claim_doc_no(row)
    reason = str(row["reason_category"] or "기타")
    reason_phrase = REASON_PHRASES.get(reason, reason)
    hours = float(row["stop_hours"] or 0)
    workers = int(row["affected_workers"] or 0)
    rate = float(row["hourly_labor_rate"] or 0)
    labor = round(hours * workers * rate, 0)
    material = float(row["material_loss_cost"] or 0)
    total = float(row["claimed_amount"] or (labor + material))
    details = (row["details"] or "").strip() or "상세 경위는 별도 협의 바랍니다."
    buyer_ceo = buyer.get("ceo_name") or "대표이사"
    buyer_name = buyer.get("company_name") or str(row["customer_name"] or "")
    seller_name = seller.get("company_name") or ""
    seller_ceo = seller.get("ceo_name") or "대표이사"
    c = CLAIM_COLS

    _paint(ws, 1, 1, 1, 5, value=seller_name, font=LETTERHEAD_FONT, align=LEFT)
    _paint(ws, 1, 6, 1, 6, value="담당", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 1, 7, 1, 7, value="팀장", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 1, 8, 1, 8, value="대표", font=BOLD_FONT, fill=HEADER_FILL)
    ws.row_dimensions[1].height = 20

    letterhead = "  ·  ".join(
        p
        for p in (
            seller.get("address", ""),
            f"TEL {seller['phone']}" if seller.get("phone") else "",
            f"FAX {seller['fax']}" if seller.get("fax") else "",
        )
        if p
    )
    _paint(ws, 2, 1, 3, 5, value=letterhead, font=SMALL_FONT, align=LEFT)
    _paint(ws, 2, 6, 3, 6, value="", fill=STAMP_FILL)
    _paint(ws, 2, 7, 3, 7, value="", fill=STAMP_FILL)
    _paint(ws, 2, 8, 3, 8, value="(인)", font=SMALL_FONT, fill=STAMP_FILL)
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 22

    _paint(ws, 4, 1, 5, c, value="업 무 협 조 전", font=TITLE_FONT)
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 18
    _paint(ws, 6, 1, 6, c, value="(손실보전금 청구)", font=SUB_FONT)
    ws.row_dimensions[6].height = 16

    _paint(ws, 7, 1, 7, 1, value="문서번호", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(ws, 7, 2, 7, 4, value=doc_no, font=LETTER_FONT, align=LEFT)
    _paint(ws, 7, 5, 7, 5, value="시행일자", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(ws, 7, 6, 7, c, value=date_text, font=LETTER_FONT)
    ws.row_dimensions[7].height = 20

    _paint(ws, 8, 1, 8, 1, value="수  신", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(
        ws, 8, 2, 8, c,
        value=f"{buyer_name}  대표이사  {buyer_ceo}  귀하",
        font=LETTER_BOLD, align=LEFT,
    )
    ws.row_dimensions[8].height = 20

    buyer_line = "  ·  ".join(
        p
        for p in (
            f"등록번호 {buyer['biz_no']}" if buyer.get("biz_no") else "",
            buyer.get("address", ""),
            f"TEL {buyer['phone']}" if buyer.get("phone") else "",
        )
        if p
    )
    _paint(ws, 9, 1, 9, 1, value="(수신처)", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(ws, 9, 2, 9, c, value=buyer_line or buyer_name, font=BODY_FONT, align=LEFT)
    ws.row_dimensions[9].height = 20

    refer = buyer.get("manager_name") or "구매팀 / 생산관리 담당"
    if buyer.get("manager_phone"):
        refer = f"{refer}  ({buyer['manager_phone']})"
    _paint(ws, 10, 1, 10, 1, value="참  조", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(ws, 10, 2, 10, 4, value=refer, font=LETTER_FONT, align=LEFT)
    _paint(ws, 10, 5, 10, 5, value="발신 담당", font=BOLD_FONT, fill=LABEL_FILL)
    sender_mgr = "  ".join(
        p for p in (seller.get("manager_name", ""), seller.get("manager_phone", "") or seller.get("phone", "")) if p
    )
    _paint(ws, 10, 6, 10, c, value=sender_mgr, font=LETTER_FONT, align=LEFT)
    ws.row_dimensions[10].height = 20

    _paint(ws, 11, 1, 11, 1, value="발  신", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(
        ws, 11, 2, 11, c,
        value=f"{seller_name}  /  {seller.get('address', '')}",
        font=LETTER_FONT, align=LEFT,
    )
    ws.row_dimensions[11].height = 20

    _paint(ws, 12, 1, 12, 1, value="제  목", font=BOLD_FONT, fill=LABEL_FILL)
    _paint(
        ws, 12, 2, 12, c,
        value=f"【손실보전금 청구 협조의 건】  -  {reason_phrase}",
        font=LETTER_BOLD, align=LEFT,
    )
    ws.row_dimensions[12].height = 22

    greeting = (
        "1. 귀사의 무궁한 발전을 기원합니다.\n\n"
        f"2. 당사는 {date_text} 귀사와의 거래 과정에서 「{reason_phrase}」 귀책으로 생산라인이 중단되어 "
        "인건비 및 자재 손실이 발생하였습니다. 아래 산출 내역과 같이 손실보전금 지급을 요청하오니 "
        "업무에 협조하여 주시기 바랍니다."
    )
    _paint(ws, 13, 1, 15, c, value=greeting, font=LETTER_FONT, align=Alignment(horizontal="left", vertical="top", wrap_text=True))
    ws.row_dimensions[13].height = 26
    ws.row_dimensions[14].height = 26
    ws.row_dimensions[15].height = 26

    _paint(ws, 16, 1, 16, c, value="3. 청구 내역", font=LETTER_BOLD, fill=LABEL_FILL, align=LEFT)
    ws.row_dimensions[16].height = 20

    _paint(ws, 17, 1, 17, 1, value="구분", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 17, 2, 17, 4, value="산출 근거", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 17, 5, 17, 5, value="수량·시간", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 17, 6, 17, 6, value="단가(원)", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 17, 7, 17, c, value="금액(원)", font=BOLD_FONT, fill=HEADER_FILL)
    ws.row_dimensions[17].height = 20

    _paint(ws, 18, 1, 18, 1, value="인건비 손실", font=BOLD_FONT)
    _paint(
        ws, 18, 2, 18, 4,
        value=f"중단시간 × 영향인원 × 시간당 인건비\n({hours:g}h × {workers}명 × {rate:,.0f}원)",
        font=BODY_FONT, align=LEFT,
    )
    _paint(ws, 18, 5, 18, 5, value=f"{hours:g}h / {workers}명", font=BODY_FONT)
    _paint(ws, 18, 6, 18, 6, value=rate, font=BODY_FONT, number_format=MONEY_FMT)
    _paint(ws, 18, 7, 18, c, value=labor, font=BOLD_FONT, number_format=MONEY_FMT)
    ws.row_dimensions[18].height = 32

    _paint(ws, 19, 1, 19, 1, value="자재 손실", font=BOLD_FONT)
    _paint(ws, 19, 2, 19, 4, value="라인 중단으로 인한 투입 자재 폐기·재작업 손실", font=BODY_FONT, align=LEFT)
    _paint(ws, 19, 5, 19, 5, value="1식", font=BODY_FONT)
    _paint(ws, 19, 6, 19, 6, value=material, font=BODY_FONT, number_format=MONEY_FMT)
    _paint(ws, 19, 7, 19, c, value=material, font=BOLD_FONT, number_format=MONEY_FMT)
    ws.row_dimensions[19].height = 22

    _paint(ws, 20, 1, 20, 4, value=f"합  계    일금  {amount_in_korean(total)}", font=LETTER_BOLD, fill=TOTAL_FILL, align=LEFT)
    _paint(ws, 20, 5, 20, 6, value="", fill=TOTAL_FILL)
    _paint(ws, 20, 7, 20, c, value=total, font=MONEY_FONT, fill=TOTAL_FILL, align=RIGHT, number_format='"₩ "#,##0')
    ws.row_dimensions[20].height = 24

    _paint(ws, 21, 1, 21, c, value="4. 발생 경위", font=LETTER_BOLD, fill=LABEL_FILL, align=LEFT)
    ws.row_dimensions[21].height = 20
    _paint(
        ws, 22, 1, 24, c,
        value=details,
        font=LETTER_FONT,
        align=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    ws.row_dimensions[22].height = 22
    ws.row_dimensions[23].height = 22
    ws.row_dimensions[24].height = 22

    bank = " ".join(
        p
        for p in (
            seller.get("bank_name", ""),
            seller.get("bank_account", ""),
            f"예금주 {seller['bank_holder']}" if seller.get("bank_holder") else "",
        )
        if p
    )
    _paint(ws, 25, 1, 25, c, value="5. 입금 계좌 및 협조 요청", font=LETTER_BOLD, fill=LABEL_FILL, align=LEFT)
    ws.row_dimensions[25].height = 20
    _paint(
        ws, 26, 1, 27, c,
        value=(
            f"입금계좌  {bank or '당사 지정 계좌'}\n"
            "위 금액을 정히 청구하오니, 확인 후 업무에 협조하여 주시기 바랍니다.  "
            "본 협조전 접수 후 7일 이내 이의가 없으면 위 내용이 확인된 것으로 합니다."
        ),
        font=LETTER_FONT,
        align=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    ws.row_dimensions[26].height = 22
    ws.row_dimensions[27].height = 22

    _paint(ws, 28, 1, 28, c, value="붙임  1. 손실 산출내역 1부.    끝.", font=LETTER_FONT, align=LEFT)
    ws.row_dimensions[28].height = 20

    _paint(ws, 29, 1, 29, 5, value=date_text, font=LETTER_BOLD)
    _paint(ws, 29, 6, 29, c, value="", fill=STAMP_FILL)
    ws.row_dimensions[29].height = 20
    _paint(ws, 30, 1, 30, 5, value=seller_name, font=LETTERHEAD_FONT)
    _paint(ws, 30, 6, 31, c, value="(인)", font=SMALL_FONT, fill=STAMP_FILL)
    ws.row_dimensions[30].height = 22
    _paint(ws, 31, 1, 31, 5, value=f"대표이사  {seller_ceo}", font=LETTER_BOLD)
    ws.row_dimensions[31].height = 22

    _paint(ws, 32, 1, 32, 1, value="수신 확인", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 32, 2, 32, 3, value="확인일자", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 32, 4, 32, 5, value="담당자", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 32, 6, 32, c, value="서명 / 날인", font=BOLD_FONT, fill=HEADER_FILL)
    _paint(ws, 33, 1, 33, 1, value=buyer_name, font=BODY_FONT)
    _paint(ws, 33, 2, 33, 3, value="", fill=WHITE)
    _paint(ws, 33, 4, 33, 5, value="", fill=WHITE)
    _paint(ws, 33, 6, 33, c, value="", fill=STAMP_FILL)
    ws.row_dimensions[32].height = 18
    ws.row_dimensions[33].height = 32

    _outline(ws, 1, 1, 33, c)
    _outline(ws, 17, 1, 20, c)
    _widths(ws, {1: 14, 2: 14, 3: 12, 4: 12, 5: 13, 6: 12, 7: 12, 8: 13})
    ws.print_area = "A1:H33"
    wb.save(path)
    return path
