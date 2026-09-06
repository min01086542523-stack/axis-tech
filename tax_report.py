"""국세청 제출용 상용직 급여대장 · 개인사업소득세 엑셀 (현장 양식)."""

from __future__ import annotations

from calendar import monthrange
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet

import hr_crypto
import hr_database as hr
from payroll_slip import COMPANY_NAME, LTC_RATE

TAX_DIR = Path(__file__).resolve().parent / "국세청신고"

GREEN = PatternFill("solid", fgColor="548235")
GREEN_HEAD = PatternFill("solid", fgColor="70AD47")
YELLOW = PatternFill("solid", fgColor="FFFF99")
WHITE = PatternFill("solid", fgColor="FFFFFF")
FONT_WHITE = Font(name="Malgun Gothic", size=9, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Malgun Gothic", size=16, bold=True)
FONT = Font(name="Malgun Gothic", size=9)
FONT_BOLD = Font(name="Malgun Gothic", size=9, bold=True)
THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONEY = "#,##0"


def _dot_date(value: str | None) -> str:
    text = (value or "").strip()[:10]
    if len(text) >= 10 and text[4] == "-":
        return text.replace("-", ".")
    return text


def _dot_ym(pay_ym: str) -> str:
    year, month = pay_ym.split("-")
    return f"{year}.{month}"


def _title_ym(pay_ym: str) -> str:
    year, month = pay_ym.split("-")
    return f"{int(year)}년 {int(month)}월"


def _work_ym(pay_ym: str) -> str:
    """지급월 기준 해당월(전월). 예: 지급 2024.10 → 해당 2024.09"""
    dt = datetime.strptime(pay_ym + "-01", "%Y-%m-%d")
    if dt.month == 1:
        return f"{dt.year - 1}.12"
    return f"{dt.year}.{dt.month - 1:02d}"


def _pay_day(pay_ym: str) -> str:
    year, month = (int(x) for x in pay_ym.split("-"))
    last = monthrange(year, month)[1]
    return f"{year}.{month:02d}.{last:02d}"


def _rrn(row: Any) -> str:
    try:
        return hr_crypto.format_rrn(hr_crypto.decrypt_rrn(row["rrn_enc"]))
    except Exception:
        return row["rrn_masked"] or ""


def _cell(ws: Worksheet, row: int, col: int, value: Any = None, *, fill=None, font=FONT, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.fill = fill or WHITE
    cell.border = THIN
    cell.alignment = CENTER
    if fmt:
        cell.number_format = fmt


def _print_setup(ws: Worksheet, last_col: int, last_row: int) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.4)
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.horizontalCentered = True


def write_regular_wage_file(path: str | Path, pay_ym: str) -> Path:
    """(주)○○ ○○○○년 ○○월 급여대장 (상용직)."""
    rows = [r for r in hr.fetch_payroll(pay_ym=pay_ym) if not hr.is_business_tax(hr.employee_tax_type(r))]
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "급여대장(상용직)"
    last_col = 24

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value=f"{COMPANY_NAME} {_title_ym(pay_ym)} 급여대장 (상용직)")
    title.font = FONT_TITLE
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    headers = (
        "NO",
        "성명",
        "사번",
        "입사일자",
        "퇴사일자",
        "주민번호",
        "기본급",
        "연차수당",
        "업무수당",
        "특근수당/잔업수당",
        "상여금",
        "지급총액",
        "건강보험",
        "요양보험",
        "국민연금",
        "고용보험",
        "연말정산 환급금",
        "소득세",
        "지방소득세",
        "공제총액",
        "차인지급액",
        "실지급액",
        "지급일",
        "비고",
    )
    for col, text in enumerate(headers, 1):
        _cell(ws, 2, col, text, fill=GREEN_HEAD, font=FONT_WHITE)
    ws.row_dimensions[2].height = 32

    money_cols = set(range(7, 23))
    if not rows:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        _cell(ws, 3, 1, "해당 월 상용직 급여 정산 내역이 없습니다.", fill=YELLOW)
        last_row = 3
    else:
        for idx, row in enumerate(rows, start=1):
            r = 2 + idx
            health = hr._pay_float(row, "health_ins")
            ltc = round(health * LTC_RATE, 0)
            ot_hol = hr._pay_float(row, "overtime") + hr._pay_float(row, "holiday_pay")
            advance = hr._pay_float(row, "other_deduction")
            remark = f"선지급 {advance:,.0f}" if advance else (row["remark"] or "")
            values = [
                idx,
                row["name"],
                row["emp_no"],
                _dot_date(row["hire_date"]),
                _dot_date(row["resign_date"] or ""),
                _rrn(row),
                hr._pay_float(row, "base_pay"),
                hr._pay_float(row, "leave_pay"),
                hr._pay_float(row, "allowance"),
                ot_hol,
                0,
                None,
                health,
                ltc,
                hr._pay_float(row, "national_pension"),
                hr._pay_float(row, "employment_ins"),
                0,
                hr._pay_float(row, "income_tax"),
                hr.payroll_local_tax(row),
                None,
                None,
                None,
                _pay_day(pay_ym),
                remark,
            ]
            for col, value in enumerate(values, 1):
                _cell(
                    ws,
                    r,
                    col,
                    value,
                    fill=YELLOW,
                    fmt=MONEY if col in money_cols else None,
                )
            _cell(ws, r, 12, f"=SUM(G{r}:K{r})", fill=YELLOW, font=FONT_BOLD, fmt=MONEY)
            _cell(ws, r, 20, f"=SUM(M{r}:S{r})", fill=YELLOW, font=FONT_BOLD, fmt=MONEY)
            _cell(ws, r, 21, f"=L{r}-T{r}", fill=YELLOW, font=FONT_BOLD, fmt=MONEY)
            _cell(ws, r, 22, f"=U{r}-{advance}", fill=YELLOW, font=FONT_BOLD, fmt=MONEY)
            ws.row_dimensions[r].height = 18
        last_data = 2 + len(rows)
        sum_row = last_data + 1
        for col in range(1, last_col + 1):
            _cell(ws, sum_row, col, fill=GREEN_HEAD, font=FONT_WHITE)
        _cell(ws, sum_row, 1, "계", fill=GREEN_HEAD, font=FONT_WHITE)
        ws.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=6)
        for col in range(7, 23):
            letter = get_column_letter(col)
            _cell(
                ws,
                sum_row,
                col,
                f"=SUM({letter}3:{letter}{last_data})",
                fill=GREEN_HEAD,
                font=FONT_WHITE,
                fmt=MONEY,
            )
        last_row = sum_row

    widths = {
        1: 6,
        2: 10,
        3: 12,
        4: 12,
        5: 12,
        6: 16,
        7: 12,
        8: 11,
        9: 11,
        10: 14,
        11: 10,
        12: 12,
        13: 11,
        14: 11,
        15: 11,
        16: 11,
        17: 13,
        18: 10,
        19: 12,
        20: 12,
        21: 12,
        22: 12,
        23: 12,
        24: 14,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"
    _print_setup(ws, last_col, last_row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_business_income_file(path: str | Path, pay_ym: str) -> Path:
    """(주)○○ ○○○○년 ○○월 개인사업소득세."""
    rows = [r for r in hr.fetch_payroll(pay_ym=pay_ym) if hr.is_business_tax(hr.employee_tax_type(r))]
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "개인사업소득세"
    last_col = 8

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value=f"{COMPANY_NAME} {_title_ym(pay_ym)} 개인사업소득세")
    title.font = FONT_TITLE
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    headers = (
        "해당월",
        "성명",
        "주민번호",
        "급여",
        "사업소득세",
        "지방소득세",
        "차인지급액",
        "지급월",
    )
    for col, text in enumerate(headers, 1):
        _cell(ws, 2, col, text, fill=GREEN_HEAD, font=FONT_WHITE)
    ws.row_dimensions[2].height = 22

    work_ym = _work_ym(pay_ym)
    pay_dot = _dot_ym(pay_ym)
    if not rows:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        _cell(ws, 3, 1, "해당 월 사업소득(3.3%) 정산 내역이 없습니다.", fill=YELLOW)
        last_row = 3
    else:
        for idx, row in enumerate(rows, start=1):
            r = 2 + idx
            gross = hr.payroll_row_gross(row)
            _cell(ws, r, 1, work_ym, fill=YELLOW)
            _cell(ws, r, 2, row["name"], fill=YELLOW)
            _cell(ws, r, 3, _rrn(row), fill=YELLOW)
            _cell(ws, r, 4, gross, fill=YELLOW, fmt=MONEY)
            _cell(ws, r, 5, f"=TRUNC(D{r}*0.03/10)*10", fill=YELLOW, fmt=MONEY)
            _cell(ws, r, 6, f"=TRUNC(E{r}*0.1/10)*10", fill=YELLOW, fmt=MONEY)
            _cell(ws, r, 7, f"=D{r}-E{r}-F{r}", fill=YELLOW, font=FONT_BOLD, fmt=MONEY)
            _cell(ws, r, 8, pay_dot, fill=YELLOW)
            ws.row_dimensions[r].height = 18
        last_data = 2 + len(rows)
        sum_row = last_data + 1
        for col in range(1, last_col + 1):
            _cell(ws, sum_row, col, fill=GREEN_HEAD, font=FONT_WHITE)
        _cell(ws, sum_row, 1, "계", fill=GREEN_HEAD, font=FONT_WHITE)
        ws.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=3)
        for col, letter in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
            _cell(
                ws,
                sum_row,
                col,
                f"=SUM({letter}3:{letter}{last_data})",
                fill=GREEN_HEAD,
                font=FONT_WHITE,
                fmt=MONEY,
            )
        last_row = sum_row

    for col, width in enumerate((12, 12, 18, 14, 14, 14, 14, 12), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"
    _print_setup(ws, last_col, last_row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def month_file_paths(pay_ym: str) -> tuple[Path, Path]:
    TAX_DIR.mkdir(parents=True, exist_ok=True)
    return (
        TAX_DIR / f"{pay_ym}_급여대장_상용직.xlsx",
        TAX_DIR / f"{pay_ym}_개인사업소득세.xlsx",
    )


def sync_hometax_month(pay_ym: str) -> tuple[Path, Path]:
    """해당 월 급여정산 자료를 현장 양식 파일에 다시 채운다."""
    regular, business = month_file_paths(pay_ym)
    write_regular_wage_file(regular, pay_ym)
    write_business_income_file(business, pay_ym)
    return regular, business
