"""AXIS TECH 로고 — 화면·인사서식 공통."""

from __future__ import annotations

from pathlib import Path
from typing import Any

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
LOGO_LIGHT = ASSETS_DIR / "axis_logo.png"
LOGO_DARK = ASSETS_DIR / "axis_logo_dark.png"
LOGO_WATERMARK = ASSETS_DIR / "axis_logo_watermark.png"

_ctk_cache: dict[tuple[int, int, int], Any] = {}

def lock_excel_pictures_no_select(xlsx_path: str | Path) -> None:
    """
    엑셀에 삽입된 이미지(로고/워터마크)가 클릭·선택되지 않게 OOXML에 noSelect를 넣는다.
    """
    import os
    import zipfile
    from xml.etree import ElementTree as ET

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        return

    xdr_ns = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"

    def _lock_drawing_xml(xml_bytes: bytes) -> bytes:
        try:
            root = ET.fromstring(xml_bytes)
        except Exception:
            return xml_bytes

        # 모든 pic에 noSelect / noMove / noResize 등을 강제한다.
        for pic in root.findall(f".//{{{xdr_ns}}}pic"):
            nv_pic_pr = pic.find(f"{{{xdr_ns}}}nvPicPr")
            if nv_pic_pr is None:
                continue
            c_nv_pic_pr = nv_pic_pr.find(f"{{{xdr_ns}}}cNvPicPr")
            if c_nv_pic_pr is None:
                continue
            pic_locks = c_nv_pic_pr.find(f"{{{a_ns}}}picLocks")
            if pic_locks is None:
                pic_locks = ET.Element(f"{{{a_ns}}}picLocks")
                c_nv_pic_pr.append(pic_locks)
            pic_locks.set("noChangeAspect", "1")
            pic_locks.set("noMove", "1")
            pic_locks.set("noResize", "1")
            pic_locks.set("noSelect", "1")

        try:
            # xml_declaration을 포함하면 일부 엑셀에서 더 안정적이다.
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        except Exception:
            return xml_bytes

    tmp = xlsx_path.with_suffix(xlsx_path.suffix + ".locktmp")
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.startswith("xl/drawings/") and info.filename.endswith(".xml"):
                    data = _lock_drawing_xml(data)
                zout.writestr(info, data)
        os.replace(tmp, xlsx_path)
    except Exception:
        try:
            if tmp.exists():
                os.remove(tmp)
        except Exception:
            pass


def logo_image(*, width: int, height: int, dark: bool = False):
    try:
        from PIL import Image
        import customtkinter as ctk
    except ImportError:
        return None
    key = (width, height, int(dark))
    cached = _ctk_cache.get(key)
    if cached is not None:
        return cached
    path = LOGO_DARK if dark else LOGO_LIGHT
    if not path.exists():
        return None
    img = Image.open(path)
    ctk_img = ctk.CTkImage(light_image=img, dark_image=img, size=(width, height))
    _ctk_cache[key] = ctk_img
    return ctk_img


def pack_logo(parent, *, width: int = 140, height: int = 142, **pack_kw):
    import customtkinter as ctk

    img = logo_image(width=width, height=height)
    if img is None:
        return None
    label = ctk.CTkLabel(parent, image=img, text="", fg_color="transparent")
    label.pack(**pack_kw)
    return label


def pack_header_logo(header, *, width: int = 58, height: int = 59) -> None:
    pack_logo(header, width=width, height=height, side="right", padx=(10, 0), pady=(0, 2))


def ensure_watermark_file(*, opacity: float = 0.28) -> Path:
    """흰 서류용으로 알파를 낮춘 흐린 로고를 만든다."""
    from PIL import Image

    key = int(round(max(0.01, min(1.0, opacity)) * 100))
    out = LOGO_WATERMARK if key == 28 else ASSETS_DIR / f"axis_logo_watermark_{key}.png"
    src_path = LOGO_DARK if LOGO_DARK.exists() else LOGO_LIGHT
    if key == 28:
        src_path = LOGO_LIGHT if LOGO_LIGHT.exists() else LOGO_DARK
    src = Image.open(src_path).convert("RGBA")
    pixels = src.load()
    width, height = src.size
    fade = max(0.01, min(1.0, opacity))
    for y in range(height):
        for x in range(width):
            _red, _green, _blue, alpha = pixels[x, y]
            if alpha == 0:
                continue
            pixels[x, y] = (160, 160, 160, max(8, int(alpha * fade)))
    src.save(out, "PNG")
    return out


def _col_px(ws, col_idx: int, last_col: int) -> float:
    from openpyxl.utils import get_column_letter

    dim = ws.column_dimensions.get(get_column_letter(col_idx))
    if dim is not None and dim.width:
        width = float(dim.width)
    elif last_col <= 4:
        width = 24.39
    elif last_col <= 6:
        width = 16.02
    else:
        width = 11.81
    return width * 7.0 + 5.0


def _row_h_px(ws, row: int) -> float:
    dim = ws.row_dimensions.get(row)
    height = float(dim.height) if dim is not None and dim.height else 15.0
    return height * 96.0 / 72.0


def _used_height_px(ws) -> float:
    last_row = int(ws.max_row or 28)
    return max(sum(_row_h_px(ws, row) for row in range(1, last_row + 1)), 360.0)


def _anchor_at(ws, x_px: float, y_px: float, width_px: float, height_px: float, last_col: int):
    """픽셀 좌표를 Excel이 따르는 OneCellAnchor로 바꾼다."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU

    remaining_x = max(0.0, x_px)
    col = 0
    for col_idx in range(1, max(1, last_col) + 1):
        width = _col_px(ws, col_idx, last_col)
        if remaining_x < width or col_idx == last_col:
            col = col_idx - 1
            break
        remaining_x -= width

    remaining_y = max(0.0, y_px)
    last_row = max(1, int(ws.max_row or 1))
    row = 0
    for row_idx in range(1, last_row + 1):
        height = _row_h_px(ws, row_idx)
        if remaining_y < height or row_idx == last_row:
            row = row_idx - 1
            break
        remaining_y -= height

    def emu(value: float) -> int:
        return int(pixels_to_EMU(max(0, value)))

    return OneCellAnchor(
        _from=AnchorMarker(
            col=col,
            colOff=emu(remaining_x),
            row=row,
            rowOff=emu(remaining_y),
        ),
        ext=XDRPositiveSize2D(emu(width_px), emu(height_px)),
    )


def write_excel_watermark(ws, last_col: int = 4, *, opacity: float = 0.15) -> None:
    """줄 간격 조정 후, 시트 정중앙에 흐린 로고를 넣는다."""
    from openpyxl.drawing.image import Image as XLImage

    if getattr(ws, "_axis_watermark", False):
        return
    if not LOGO_LIGHT.exists() and not LOGO_DARK.exists():
        return
    path = ensure_watermark_file(opacity=opacity)
    img = XLImage(str(path))
    target = 200
    ratio = target / float(img.width or 1)
    img.width = target
    img.height = max(1, int((img.height or target) * ratio))
    sheet_w = sum(_col_px(ws, col, last_col) for col in range(1, last_col + 1))
    sheet_h = _used_height_px(ws)
    x = max(0, (sheet_w - img.width) / 2)
    y = max(0, (sheet_h - img.height) / 2)
    img.anchor = _anchor_at(ws, x, y, img.width, img.height, last_col)
    ws.add_image(img)
    ws._axis_watermark = True


def write_excel_logo(ws, row: int, last_col: int = 4, *, width_px: int = 88) -> int:
    """「끝」 줄의 맨 오른쪽 칸에 로고를 붙인다. 셀 앵커라 Excel이 위치를 지킨다."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    if not LOGO_DARK.exists():
        return row + 1
    if getattr(ws, "_axis_footer_logo", False):
        return row + 1
    img = XLImage(str(LOGO_DARK))
    ratio = width_px / float(img.width or 1)
    img.width = width_px
    img.height = max(1, int((img.height or width_px) * ratio))
    needed_pt = img.height * 72.0 / 96.0
    ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 0), needed_pt)
    for col_idx in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col_idx)
        if cell.alignment is None or cell.alignment.horizontal is None:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    img.anchor = f"{get_column_letter(last_col)}{row}"
    ws.add_image(img)
    ws._axis_footer_logo = True
    return row + 1


def _end_mark_row(ws) -> int:
    last_row = int(ws.max_row or 1)
    for row in range(last_row, 0, -1):
        value = ws.cell(row=row, column=1).value
        if isinstance(value, str) and "끝" in value:
            return row
    return last_row


def stamp_workbook_logos(workbook) -> None:
    from openpyxl.utils import get_column_letter

    for ws in workbook.worksheets:
        last_col = int(ws.max_column or 4)
        if not getattr(ws, "_axis_footer_logo", False):
            write_excel_logo(ws, _end_mark_row(ws), last_col)
        write_excel_watermark(ws, last_col, opacity=0.15)
        last_col = int(ws.max_column or 4)
        last_row = int(ws.max_row or 1)
        ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    try:
        import excel_export

        excel_export.normalize_workbook(workbook)
    except Exception:
        pass


def stamp_workbook_watermarks(workbook, *, opacity: float = 0.15) -> None:
    for ws in workbook.worksheets:
        last_col = int(ws.max_column or 4)
        write_excel_watermark(ws, last_col, opacity=opacity)
    try:
        import excel_export

        excel_export.normalize_workbook(workbook)
    except Exception:
        pass
