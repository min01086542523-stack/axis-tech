"""목록 화면용 엑셀 내보내기."""

from __future__ import annotations

import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk

OUTPUT_DIR = Path(__file__).resolve().parent / "서식출력"
_CREATE_NO_WINDOW = 0x08000000
_XL_NORMAL_VIEW = 1
_XL_MAXIMIZED = -4137


def default_export_path(default_name: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    name = Path(default_name).name
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    stem = "".join(ch for ch in Path(name).stem if ch not in r'\/:*?"<>|') or "내보내기"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return OUTPUT_DIR / f"{stem}_{stamp}.xlsx"


def normalize_workbook(workbook) -> None:
    """페이지 레이아웃·분할 창이 아니라 일반 보기 한 창으로 저장한다."""
    views = list(getattr(workbook, "views", None) or [])
    if views:
        workbook.views = [views[0]]
        view = workbook.views[0]
        try:
            view.minimized = False
            view.visibility = "visible"
            view.windowWidth = 28800
            view.windowHeight = 17400
            view.xWindow = 120
            view.yWindow = 120
        except Exception:
            pass
    for ws in workbook.worksheets:
        try:
            sheet_view = ws.sheet_view
            sheet_view.view = "normal"
            if not getattr(sheet_view, "zoomScale", None):
                sheet_view.zoomScale = 100
            sheet_view.zoomScaleNormal = 100
            pane = getattr(sheet_view, "pane", None)
            state = str(getattr(pane, "state", "") or "").lower() if pane is not None else ""
            if state == "split":
                sheet_view.pane = None
                ws.freeze_panes = None
        except Exception:
            pass


def open_exported(path: str | Path) -> None:
    target = Path(path)
    if not target.exists():
        raise FileNotFoundError(str(target))
    full = os.path.normpath(str(target.resolve()))
    if target.suffix.lower() in {".xlsx", ".xlsm"}:
        _launch_excel(full)
        return
    os.startfile(full)


def _launch_excel(full: str) -> None:
    if sys.platform.startswith("win") and _open_excel_com(full):
        return
    os.startfile(full)
    if sys.platform.startswith("win"):
        _fix_excel_layout_async(full)


def _open_excel_com(full: str) -> bool:
    try:
        import win32com.client
    except ImportError:
        return False
    try:
        try:
            excel = win32com.client.GetActiveObject("Excel.Application")
        except Exception:
            excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = True
        book = None
        wanted = os.path.normpath(full).lower()
        for opened in excel.Workbooks:
            try:
                if os.path.normpath(str(opened.FullName)).lower() == wanted:
                    book = opened
                    break
            except Exception:
                continue
        if book is None:
            book = excel.Workbooks.Open(full)
        _arrange_excel(excel, book)
        return True
    except Exception:
        return False


def _arrange_excel(excel, book) -> None:
    try:
        excel.Windows.BreakSideBySide()
    except Exception:
        pass
    try:
        excel.WindowState = _XL_MAXIMIZED
    except Exception:
        pass
    try:
        while book.Windows.Count > 1:
            book.Windows(book.Windows.Count).Close()
    except Exception:
        pass
    try:
        window = book.Windows(1)
        window.View = _XL_NORMAL_VIEW
        if window.Split:
            window.Split = False
        window.WindowState = _XL_MAXIMIZED
        window.Activate()
    except Exception:
        pass
    try:
        book.Activate()
    except Exception:
        pass


def _fix_excel_layout_async(full: str) -> None:
    script = (
        "$ErrorActionPreference = 'SilentlyContinue';"
        "$path = $env:MES_EXCEL_PATH;"
        "for ($i = 0; $i -lt 24; $i++) {"
        "  Start-Sleep -Milliseconds 250;"
        "  try { $excel = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application') } catch { continue };"
        "  if (-not $excel) { continue };"
        "  $excel.Windows.BreakSideBySide() | Out-Null;"
        "  $excel.WindowState = -4137;"
        "  foreach ($wb in @($excel.Workbooks)) {"
        "    $name = [string]$wb.FullName;"
        "    if ($path -and $name -and ($name.ToLower() -ne $path.ToLower())) { continue };"
        "    while ($wb.Windows.Count -gt 1) { $wb.Windows.Item($wb.Windows.Count).Close() };"
        "    $w = $wb.Windows.Item(1);"
        "    $w.View = 1;"
        "    if ($w.Split) { $w.Split = $false };"
        "    $w.WindowState = -4137;"
        "  };"
        "  break;"
        "}"
    )
    env = os.environ.copy()
    env["MES_EXCEL_PATH"] = full
    try:
        subprocess.Popen(
            ["powershell", "-NoProfile", "-WindowStyle", "Hidden", "-Command", script],
            env=env,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=_CREATE_NO_WINDOW,
        )
    except Exception:
        pass


def export_sheets_to_xlsx(parent, default_name: str, sheets: tuple) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        messagebox.showerror(
            "엑셀 내보내기",
            "openpyxl 패키지가 필요합니다.\n터미널에서 pip install openpyxl 을 실행하세요.",
            parent=parent,
        )
        return

    if not any(sheet[2] for sheet in sheets):
        messagebox.showwarning("엑셀 내보내기", "내보낼 데이터가 없습니다.", parent=parent)
        return

    path = default_export_path(default_name)
    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6AA5")
    header_align = Alignment(horizontal="center", vertical="center")

    for index, (title, headers, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title[:31]
        for col_idx, heading in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=heading)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row_idx, values in enumerate(rows, start=2):
            for col_idx, value in enumerate(values, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=value)
        for col_idx, heading in enumerate(headers, start=1):
            max_len = len(str(heading))
            for row_idx in range(2, len(rows) + 2):
                max_len = max(max_len, len(str(sheet.cell(row=row_idx, column=col_idx).value or "")))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)
        for col_idx in range(1, 16):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 13
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"

    try:
        import brand

        brand.stamp_workbook_watermarks(workbook)
        normalize_workbook(workbook)
        workbook.save(path)
        open_exported(path)
    except OSError as exc:
        messagebox.showwarning(
            "열기",
            f"서식은 저장했습니다.\n{path}\n\n파일을 열지 못했습니다: {exc}",
            parent=parent,
        )


def export_tree_to_xlsx(
    tree: ttk.Treeview,
    parent,
    default_name: str,
    numeric_columns: set[str] | None = None,
) -> None:
    numeric_columns = numeric_columns or set()
    columns = list(tree["columns"])
    headers = tuple(tree.heading(col)["text"] for col in columns)
    data = []
    for item_id in tree.get_children():
        values = tree.item(item_id, "values")
        row = []
        for col, raw in zip(columns, values):
            value: object = raw
            if col in numeric_columns:
                cleaned = str(raw).replace(",", "").strip()
                try:
                    if cleaned == "":
                        value = 0
                    elif "." in cleaned:
                        value = float(cleaned)
                    else:
                        value = int(cleaned)
                except ValueError:
                    value = raw
            row.append(value)
        data.append(tuple(row))
    export_sheets_to_xlsx(
        parent=parent,
        default_name=default_name,
        sheets=(("내보내기", headers, tuple(data)),),
    )
